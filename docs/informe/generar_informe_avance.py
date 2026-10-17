#!/usr/bin/env python3
"""Genera el informe Word de la presentación del avance.

    python docs/informe/generar_informe_avance.py

Formato pedido: carátula, índice de contenido, índice de imágenes, Times New
Roman 12, interlineado doble, espaciado entre párrafos 0, número de página en
el pie, encabezado con el logo institucional y los apellidos y nombres, citas y
referencias en APA 7 (normas ISO y legislación peruana).

Los índices son campos de Word: Word los actualiza al abrir el documento
(se pide con `updateFields`). Si se edita el texto, actualizarlos con F9. La
versión entregada en docs/informe/ se pasó por LibreOffice para que los índices
ya vengan rellenos con sus números de página (y se exportó a PDF).
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
FIG = HERE / "figuras"
OUT = HERE / "Informe de avance - Video en vivo, Audio e Imagen.docx"

# Apellidos y nombres que van en el ENCABEZADO de cada página. Si el informe
# es individual, deja solo el del estudiante que lo presenta.
ESTUDIANTES = [
    "Hermenegildo Rumiche, Andy Humberto",
    "Mendoza Avila, Jorge Luis",
    "Lopez Gonzalez, Jorge",
    "Coronel Ruiz, Brayan",
    "Cortez Acon, Jonaiker Mirosevic",
    "Sifuentes Berrocal, Edy Anthony",
    "Pereda Obando, Luis Go",
]
DOCENTE = "Mendoza Corpus, Carlos Alfredo"
FONT = "Times New Roman"


# --------------------------------------------------------------------------- #
# utilidades de formato
# --------------------------------------------------------------------------- #
def set_font(run, size=12, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(k), FONT)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def fmt(par, *, align=None, indent=True, spacing="double", keep_next=False):
    pf = par.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = (WD_LINE_SPACING.DOUBLE if spacing == "double"
                            else WD_LINE_SPACING.SINGLE)
    pf.first_line_indent = Cm(1.27) if indent else Cm(0)
    if align is not None:
        par.alignment = align
    if keep_next:
        pf.keep_with_next = True
    return par


def add_field(par, instr: str, placeholder: str = ""):
    """Campo de Word (PAGE, TOC, SEQ...)."""
    def fld(kind):
        r = OxmlElement("w:r")
        c = OxmlElement("w:fldChar")
        c.set(qn("w:fldCharType"), kind)
        if kind == "begin":
            c.set(qn("w:dirty"), "true")
        r.append(c)
        return r
    p = par._p
    p.append(fld("begin"))
    r = OxmlElement("w:r")
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = f" {instr} "
    r.append(t)
    p.append(r)
    p.append(fld("separate"))
    run = par.add_run(placeholder)
    set_font(run)
    p.append(fld("end"))
    return run


def text(doc, parts, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True):
    """Párrafo de cuerpo. `parts`: str o lista de (texto, {'b':..,'i':..})."""
    par = fmt(doc.add_paragraph(), align=align, indent=indent)
    if isinstance(parts, str):
        parts = [(parts, {})]
    for t, st in parts:
        set_font(par.add_run(t), bold=st.get("b"), italic=st.get("i"))
    return par


def heading(doc, title, level):
    """Títulos APA 7: nivel 1 centrado y negrita; 2 izquierda negrita;
    3 izquierda negrita cursiva. Usan los estilos Heading para el índice."""
    par = doc.add_heading(level=level)
    fmt(par, indent=False, keep_next=True,
        align=WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT)
    set_font(par.add_run(title), bold=True, italic=(level == 3))
    return par


def bullets(doc, items):
    for it in items:
        par = fmt(doc.add_paragraph(style="List Bullet"),
                  align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=False)
        if isinstance(it, str):
            it = [(it, {})]
        for t, st in it:
            set_font(par.add_run(t), bold=st.get("b"), italic=st.get("i"))


def page_break(doc):
    par = fmt(doc.add_paragraph(), indent=False)
    par.add_run().add_break(WD_BREAK.PAGE)


FIG_N = [0]
TAB_N = [0]


def figure(doc, path, title, note, width_cm=15.5):
    """Figura APA 7: «Figura N» en negrita y el título en cursiva encima; la
    nota debajo. Número y título van en un mismo párrafo (con salto de línea)
    para que el índice de imágenes muestre el título completo."""
    FIG_N[0] += 1
    cap = fmt(doc.add_paragraph(style="Caption"), indent=False, keep_next=True,
              align=WD_ALIGN_PARAGRAPH.LEFT)
    set_font(cap.add_run("Figura "), bold=True, italic=False)
    r = add_field(cap, "SEQ Figura \\* ARABIC", str(FIG_N[0]))
    r.bold = True
    br = cap.add_run()
    br.add_break(WD_BREAK.LINE)
    set_font(br)
    set_font(cap.add_run(title), italic=True, bold=False)
    for run in cap.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    pic = fmt(doc.add_paragraph(), indent=False, keep_next=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, spacing="single")
    pic.add_run().add_picture(str(path), width=Cm(width_cm))
    n = fmt(doc.add_paragraph(), indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_font(n.add_run("Nota. "), italic=True)
    set_font(n.add_run(note))


def table(doc, title, header, rows, widths_cm, note=None):
    """Tabla APA 7 (solo líneas horizontales)."""
    TAB_N[0] += 1
    cap = fmt(doc.add_paragraph(), indent=False, keep_next=True)
    set_font(cap.add_run(f"Tabla {TAB_N[0]}"), bold=True)
    cap2 = fmt(doc.add_paragraph(), indent=False, keep_next=True)
    set_font(cap2.add_run(title), italic=True)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate([header] + rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Cm(widths_cm[j])
            par = fmt(cell.paragraphs[0], indent=False, spacing="single",
                      align=WD_ALIGN_PARAGRAPH.LEFT)
            set_font(par.add_run(val), size=11, bold=(i == 0))
    _apa_borders(t, len(rows))
    if note:
        n = fmt(doc.add_paragraph(), indent=False)
        set_font(n.add_run("Nota. "), italic=True)
        set_font(n.add_run(note))
    else:
        fmt(doc.add_paragraph(), indent=False)


def _apa_borders(t, nrows):
    tbl = t._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, val in (("top", "single"), ("bottom", "single"), ("left", "nil"),
                      ("right", "nil"), ("insideH", "nil"), ("insideV", "nil")):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), val)
        if val == "single":
            e.set(qn("w:sz"), "8")
        borders.append(e)
    tblPr.append(borders)
    for cell in t.rows[0].cells:                 # línea bajo el encabezado
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single")
        bot.set(qn("w:sz"), "8")
        b.append(bot)
        tcPr.append(b)
    # repetir encabezado si la tabla cambia de página
    trPr = t.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    trPr.append(h)


def reference(doc, parts):
    """Referencia APA 7: sangría francesa de 1.27 cm."""
    par = fmt(doc.add_paragraph(), indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    par.paragraph_format.left_indent = Cm(1.27)
    par.paragraph_format.first_line_indent = Cm(-1.27)
    for t, st in parts:
        set_font(par.add_run(t), italic=st.get("i"))


# --------------------------------------------------------------------------- #
def setup(doc):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)        # A4
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(1.0)

    for name in ("Normal", "Caption", "List Bullet", "Heading 1", "Heading 2",
                 "Heading 3", "Title"):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(12)
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(k), FONT)
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if fonts.get(qn(k)) is not None:
                del fonts.attrib[qn(k)]
        pf = st.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.bold = True
    # Estilos de los índices: también Times 12 a doble espacio
    for name in ("TOC 1", "TOC 2", "TOC 3", "Table of Figures"):
        try:
            st = doc.styles[name]
        except KeyError:
            from docx.enum.style import WD_STYLE_TYPE
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(12)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        if name == "TOC 2":
            st.paragraph_format.left_indent = Cm(0.6)
        if name == "TOC 3":
            st.paragraph_format.left_indent = Cm(1.2)

    # Word actualiza índices y números de figura al abrir el documento
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)

    # Encabezado: logo + apellidos y nombres
    hdr = sec.header
    hdr.is_linked_to_previous = False
    ht = hdr.add_table(rows=1, cols=2, width=Cm(16))
    ht.alignment = WD_TABLE_ALIGNMENT.CENTER
    ht.autofit = False
    c0, c1 = ht.rows[0].cells
    for col, w in zip(ht.columns, (Cm(1.8), Cm(14.2))):
        col.width = w
    c0.width, c1.width = Cm(1.8), Cm(14.2)
    grid = ht._tbl.tblGrid.findall(qn("w:gridCol"))
    for gc, w in zip(grid, (Cm(1.8), Cm(14.2))):
        gc.set(qn("w:w"), str(int(w.twips)))
    lp = fmt(c0.paragraphs[0], indent=False, spacing="single")
    lp.add_run().add_picture(str(FIG / "logo_upao.png"), height=Cm(1.6))
    first = True
    for i in range(0, len(ESTUDIANTES), 3):
        par = c1.paragraphs[0] if first else c1.add_paragraph()
        first = False
        fmt(par, indent=False, spacing="single", align=WD_ALIGN_PARAGRAPH.RIGHT)
        set_font(par.add_run(" · ".join(ESTUDIANTES[i:i + 3])), size=8)
    hp = hdr.paragraphs[0]
    fmt(hp, indent=False, spacing="single")
    _bottom_border(hp)

    # Pie: número de página
    fp = fmt(sec.footer.paragraphs[0], indent=False, spacing="single",
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_field(fp, "PAGE", "1")


def _bottom_border(par):
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3864")):
        bot.set(qn(k), v)
    b.append(bot)
    pPr.append(b)


# --------------------------------------------------------------------------- #
def caratula(doc):
    c = WD_ALIGN_PARAGRAPH.CENTER
    for t, b in (("UNIVERSIDAD PRIVADA ANTENOR ORREGO", True),
                 ("FACULTAD DE INGENIERÍA", True),
                 ("ESCUELA PROFESIONAL DE INGENIERÍA DE SISTEMAS E INTELIGENCIA ARTIFICIAL", True)):
        set_font(fmt(doc.add_paragraph(), align=c, indent=False).add_run(t), bold=b)
    p = fmt(doc.add_paragraph(), align=c, indent=False, spacing="single")
    p.add_run().add_picture(str(FIG / "logo_upao.png"), height=Cm(3.6))
    set_font(fmt(doc.add_paragraph(), align=c, indent=False).add_run(
        "Informe de avance"), bold=True)
    set_font(fmt(doc.add_paragraph(), align=c, indent=False).add_run(
        "Sistema inteligente para la prevención de caídas y riesgos en escaleras: "
        "operación exclusiva en vídeo en vivo, integración de los módulos de audio "
        "e imagen y compatibilidad con GPU GTX 1650"), bold=True)
    set_font(fmt(doc.add_paragraph(), align=c, indent=False).add_run("Estudiantes:"), bold=True)
    for e in ESTUDIANTES:
        set_font(fmt(doc.add_paragraph(), align=c, indent=False, spacing="single").add_run(e))
    set_font(fmt(doc.add_paragraph(), align=c, indent=False).add_run("Docente:"), bold=True)
    set_font(fmt(doc.add_paragraph(), align=c, indent=False, spacing="single").add_run(DOCENTE))
    set_font(fmt(doc.add_paragraph(), align=c, indent=False).add_run("Curso:"), bold=True)
    set_font(fmt(doc.add_paragraph(), align=c, indent=False, spacing="single").add_run(
        "Percepción Computacional"))
    set_font(fmt(doc.add_paragraph(), align=c, indent=False).add_run("Trujillo – Perú, 2026"))
    page_break(doc)


def indices(doc):
    heading(doc, "Índice de contenido", 1).style = doc.styles["Title"]
    p = fmt(doc.add_paragraph(), indent=False)
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u',
              "Pulse F9 (o abra el documento en Word) para actualizar el índice.")
    page_break(doc)
    heading(doc, "Índice de imágenes", 1).style = doc.styles["Title"]
    p = fmt(doc.add_paragraph(), indent=False)
    add_field(p, 'TOC \\h \\z \\c "Figura"',
              "Pulse F9 (o abra el documento en Word) para actualizar el índice.")
    page_break(doc)


def _title_style(doc):
    """El estilo Title se usa para los encabezados de los índices: igual que un
    título de nivel 1 pero sin entrar en el índice de contenido."""
    st = doc.styles["Title"]
    st.font.bold = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = st.element.get_or_add_pPr()
    for b in pPr.findall(qn("w:pBdr")):
        pPr.remove(b)


# --------------------------------------------------------------------------- #
def cuerpo(doc):
    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    heading(doc, "1. Introducción", 1)
    text(doc, "El presente informe documenta las actividades desarrolladas para la "
         "presentación del avance del Sistema inteligente para la prevención de caídas y "
         "riesgos en escaleras de la Universidad Privada Antenor Orrego. El sistema vigila "
         "una escalera con una cámara fija y, mediante visión por computadora (YOLOv8 y "
         "MediaPipe), avisa al personal de seguridad cuando detecta caídas, pérdida de "
         "equilibrio, tránsito sin pasamanos, distracción con el celular u obstáculos en "
         "los escalones.")
    text(doc, "En esta etapa el cliente solicitó cinco cambios: (a) que el sistema funcione "
         "en una tarjeta gráfica NVIDIA GeForce GTX 1650; (b) eliminar todo lo relativo al "
         "vídeo subido y dejar únicamente el vídeo en vivo; (c) integrar, en una carpeta "
         "aparte y sin interferir con el vídeo, la función de audio desarrollada por otro "
         "grupo; (d) integrar el módulo de imagen de un tercer grupo en sustitución de la "
         "imagen existente; y (e) añadir una función interna para detener el proyecto. "
         "Como restricción expresa, ningún modelo debía reentrenarse.")
    text(doc, "El trabajo se orientó por el modelo de calidad del producto de la norma "
         "ISO/IEC 25010 (International Organization for Standardization [ISO], 2023c), "
         "por la gestión de sistemas y riesgos de inteligencia artificial (ISO, 2023a, "
         "2023b) y por la Ley N.° 29733, Ley de Protección de Datos Personales "
         "(Congreso de la República del Perú, 2011), dado que el sistema procesa imágenes "
         "y sonidos de personas.")

    heading(doc, "2. Objetivos", 1)
    heading(doc, "2.1. Objetivo general", 2)
    text(doc, "Adaptar el sistema de prevención de caídas en escaleras para que opere "
         "exclusivamente con vídeo en vivo en equipos con una GPU GTX 1650, integrando de "
         "forma modular los trabajos de audio e imagen de los otros grupos del curso, sin "
         "reentrenar modelos y sin degradar la detección visual.")
    heading(doc, "2.2. Objetivos específicos", 2)
    bullets(doc, [
        "Retirar las rutas, la interfaz y los modelos asociados a imágenes y vídeos subidos.",
        "Garantizar la ejecución en una GPU de 4 GB de VRAM de arquitectura Turing (sm_75).",
        "Incorporar el módulo de percepción acústica con todas sus funciones activas.",
        "Incorporar el módulo de análisis de imagen en reemplazo de la pestaña anterior.",
        "Implementar una parada ordenada del sistema que libere cámara, micrófono y modelos.",
        "Verificar los cambios con pruebas automatizadas y documentarlos.",
    ])

    heading(doc, "3. Marco normativo", 1)
    text(doc, "Las normas ISO son de adopción voluntaria; en este proyecto son de "
         "cumplimiento obligatorio porque el documento de especificación de requisitos "
         "del curso las exige. La especificación sigue la norma de ingeniería de "
         "requisitos ISO/IEC/IEEE 29148 (ISO, 2018a), y cada cambio de esta etapa se "
         "trazó a una característica de calidad de ISO/IEC 25010 (ISO, 2023c):")
    bullets(doc, [
        [("Adecuación funcional y compatibilidad: ", {"b": True}),
         ("los módulos de audio e imagen coexisten con el vídeo en el mismo panel sin "
          "compartir su hilo de inferencia (ISO, 2023c).", {})],
        [("Eficiencia de desempeño y capacidad: ", {"b": True}),
         ("perfil de hardware para 4 GB de VRAM y carga diferida de modelos (ISO, 2023c).", {})],
        [("Mantenibilidad (modularidad): ", {"b": True}),
         ("cada módulo vive en su carpeta y se carga de forma aislada (ISO, 2023c).", {})],
        [("Fiabilidad: ", {"b": True}),
         ("si un módulo falla, el vídeo en vivo continúa; la parada es ordenada (ISO, 2023c).", {})],
    ])
    text(doc, "Respecto a la inteligencia artificial, la norma ISO/IEC 42001 exige "
         "trazabilidad de los modelos y de sus decisiones (ISO, 2023b); por ello los "
         "pesos de los modelos se verifican con su suma SHA-256 y las alertas se registran "
         "en archivos de auditoría. La norma ISO/IEC 23894 orienta la gestión de riesgos "
         "de la IA (ISO, 2023a): las alertas acústicas se presentan como «posible caída» "
         "que requiere verificación humana, y el módulo de imagen nunca emite una caída "
         "confirmada a partir de una sola fotografía. La terminología empleada sigue "
         "ISO/IEC 22989 (ISO, 2022b).")
    text(doc, "En seguridad de la información, el panel conserva la autenticación por "
         "token y la nueva función de parada solo se acepta desde el propio equipo cuando "
         "no hay token configurado, en línea con los controles de acceso de ISO/IEC 27001 "
         "(ISO, 2022a). Como el objetivo del sistema es reducir caídas, se relaciona también "
         "con la gestión de la seguridad y salud en el trabajo de ISO 45001 (ISO, 2018b).")
    text(doc, "En materia legal, las imágenes y la voz que permiten identificar a una "
         "persona son datos personales según la Ley N.° 29733 (Congreso de la República "
         "del Perú, 2011) y su reglamento (Ministerio de Justicia y Derechos Humanos "
         "[MINJUSDH], 2024). El tratamiento mediante cámaras se rige además por la "
         "Directiva N.° 01-2020-JUS/DGTAIPD (Dirección General de Transparencia, Acceso a "
         "la Información Pública y Protección de Datos Personales, 2020). En "
         "consecuencia: los rostros de las evidencias se pixelan; el módulo de audio "
         "procesa el sonido en memoria y solo guarda metadatos de las alertas, nunca "
         "grabaciones; y no se versionan imágenes ni audios reales de personas. "
         "Finalmente, la Ley N.° 31814 promueve el uso de la inteligencia artificial "
         "de manera responsable, segura y ética (Congreso de la República del Perú, 2023), "
         "principio que se refleja en la supervisión humana de todas las alertas.")

    heading(doc, "4. Actividades desarrolladas", 1)
    heading(doc, "4.1. Diagnóstico del proyecto y de los repositorios", 2)
    text(doc, "Se revisó la rama expo del proyecto, que tenía cuatro pestañas (Imagen, "
         "Vídeo, Cámara en vivo y Clases), y los dos repositorios a integrar: "
         "brigitteve/percepcion (audio) y AndreCastAst/Repo-video-imagen-audio, rama "
         "s7-imagen (imagen). Se identificó que el repositorio de audio contenía también "
         "una versión antigua del vídeo (su propio servidor y motor de riesgo), que por "
         "indicación del cliente no se migró; solo se tomó su parte acústica. En el "
         "repositorio de imagen se identificó que su ruta de análisis reutilizaba el "
         "motor de vídeo y reiniciaba su seguimiento, por lo que se rediseñó para no "
         "interferir con el vídeo en vivo.")
    figure(doc, FIG / "fig_arquitectura.png",
           "Arquitectura del sistema tras la integración",
           "Capas del sistema: entradas (vídeo en vivo, micrófono y fotografía), panel web, "
           "API Flask, servicios, módulos aparte (audio e imagen), núcleo de percepción y "
           "salidas. Elaboración propia (tools/render_architecture.py).")

    heading(doc, "4.2. Operación exclusiva con vídeo en vivo", 2)
    text(doc, "Se eliminaron las rutas /api/detect (imagen subida) y /api/video (vídeo "
         "subido), la recodificación H.264 que solo servía a esos vídeos, la mejora de "
         "calidad previa a la detección de archivos subidos y el modelo grande "
         "RealESRGAN_x4plus (16,7 millones de parámetros), utilizado únicamente por esas "
         "subidas. En el panel se retiraron las pestañas Imagen y Vídeo. Se conservaron "
         "la cámara del navegador (/api/frame) y la cámara IP o USB del servidor "
         "(/api/stream), ambas en vivo, junto con la mejora de calidad Real-ESRGAN en "
         "tiempo real.")
    figure(doc, FIG / "fig_camara_en_vivo.png",
           "Pestaña «Cámara en vivo» del panel de supervisión",
           "Captura con la cámara simulada del navegador. Se observan el botón «Detener "
           "sistema», el estado de los módulos de audio e imagen y el aviso de que el "
           "análisis acústico puede funcionar a la vez que la cámara.")

    heading(doc, "4.3. Compatibilidad con la GPU GTX 1650", 2)
    text(doc, "La GTX 1650 es de arquitectura Turing (sm_75), con 4 GB de VRAM y sin "
         "núcleos Tensor. Se comprobó que la versión de PyTorch del proyecto "
         "(2.11.0+cu128) ya incluye núcleos para sm_75, por lo que no se requirió otra "
         "instalación. Se implementó un perfil de hardware automático (baja_vram) que se "
         "activa cuando la GPU tiene 4,5 GB o menos o pertenece a la serie GTX 16: "
         "desactiva la media precisión (FP16) en Real-ESRGAN, porque estas tarjetas "
         "presentan fallos conocidos con FP16, y reduce la entrada de la mejora en vivo "
         "de 480 a 320 píxeles. La detección no se modificó (sigue a 640 píxeles) para "
         "no alterar la precisión del modelo, que no se reentrena.")
    table(doc, "Medición con el perfil de la GTX 1650",
          ["Medida", "Resultado"],
          [["Memoria de GPU disponible en la prueba", "Limitada a 3,6 GB (equivalente a una GTX 1650)"],
           ["Pico de VRAM con vídeo, mejora en vivo e imagen", "0,32 GB"],
           ["Real-ESRGAN en FP32, entrada de 320 px", "≈ 19 ms por cuadro"],
           ["Audio (reglas + clasificador ML), bloque de 500 ms", "≈ 165 ms de CPU"],
           ["Modelos de imagen (YOLO11)", "Se cargan solo al primer uso"]],
          [8.5, 7.5],
          "Medido en una RTX 5070 Laptop forzando el perfil baja_vram y limitando la memoria "
          "con torch.cuda.set_per_process_memory_fraction. Los tiempos en una GTX 1650 real "
          "serán mayores; la vista en vivo se mantiene fluida porque la detección es asíncrona.")

    heading(doc, "4.4. Integración del módulo de audio", 2)
    text(doc, "El código acústico se trasladó a la carpeta modulos/audio, separada del "
         "vídeo. Se conservaron todas sus funciones: captura del micrófono en el navegador "
         "(AudioWorklet, 16 kHz, bloques de 500 ms), detección de pisadas y cadencia "
         "(caminando o corriendo), impacto súbito de caída (pico RMS y factor de cresta), "
         "grito o pedido de auxilio (frecuencia fundamental con pYIN), sensibilidad alta, "
         "media o baja, demostración con la sesión simulada del cuaderno original, "
         "análisis de archivos de audio, pitido y notificaciones. Se creó la pestaña y "
         "el botón «Audio» del panel.")
    text(doc, "El grupo había dejado su clasificador de aprendizaje automático (un "
         "perceptrón multicapa calibrado de seis clases) solo en scripts; para que todas "
         "sus funciones estuvieran activas, se ejecuta en paralelo al detector por "
         "reglas sobre el mismo audio. El modelo se usa tal cual, sin reentrenar: su "
         "suma SHA-256 coincide con la del repositorio original. Como el grupo fijó "
         "scikit-learn 1.9.1, que exige Python 3.11, se utilizó la versión 1.7.2 y se "
         "verificó con sus propias pruebas que los pesos y las predicciones son "
         "idénticos. Las alertas acústicas se guardan como metadatos y se muestran "
         "también en el panel de alertas de la cámara, en la descarga ZIP y en el "
         "borrado, conforme a la integración realizada por ese grupo. El análisis de "
         "archivos acepta solo audio, de modo que el sistema no recibe vídeos.")
    figure(doc, FIG / "fig_modulo_audio.png",
           "Pestaña «Audio»: módulo de percepción acústica",
           "Micrófono simulado del navegador y resultado de la sesión simulada de "
           "final.py: caminando, corriendo, caída e impacto y pedido de auxilio, con sus "
           "niveles de riesgo. A la derecha, las alertas acústicas y los criterios.")

    heading(doc, "4.5. Integración del módulo de imagen", 2)
    text(doc, "El módulo del grupo de imagen se trasladó a la carpeta modulos/imagen y "
         "reemplazó a la pestaña de imagen anterior. Incluye la medición de calidad "
         "óptica (luminosidad, contraste, nitidez, sombras y ruido), el preprocesamiento "
         "adaptativo (original, adaptativo, CLAHE y gamma), la detección de personas y "
         "celulares con YOLO11n, la estimación de postura con YOLO11n-Pose, la asociación "
         "entre pose y persona, los estados posturales «para revisión», el análisis de "
         "fotogramas de la cámara y el vector de características imagen-0.1.0. El único "
         "cambio a su código fue cargar los modelos en el primer uso, desde la carpeta "
         "models/imagen, para no ocupar memoria de la GPU mientras no se utilizan. Su "
         "ruta que pasaba la fotografía por el motor de vídeo se sustituyó por una de "
         "solo lectura que obtiene las características del último fotograma en vivo sin "
         "alterar el vídeo.")
    figure(doc, FIG / "fig_modulo_imagen.png",
           "Pestaña «Imagen»: análisis de una fotografía con YOLO11 y YOLO11-Pose",
           "Imagen de ejemplo de Ultralytics (bus.jpg). Se muestran las personas, sus "
           "articulaciones, la calidad óptica, el preprocesamiento aplicado, las "
           "advertencias y el JSON del módulo. La franja roja superior es una alerta "
           "acústica de la prueba, visible en cualquier pestaña.")

    heading(doc, "4.6. Función interna para detener el sistema", 2)
    text(doc, "Se implementó un servicio de ciclo de vida que detiene el proyecto de forma "
         "ordenada. Puede invocarse con el botón «Detener sistema» del panel, con el "
         "comando python -m backend --stop o con Ctrl+C en la consola. Primero deja de "
         "aceptar trabajo nuevo; después cierra las sesiones de audio, libera los modelos "
         "de imagen, detiene la mejora en vivo y el hilo de inferencia (liberando la "
         "memoria de la GPU) y, por último, apaga el servidor web. En el navegador, el "
         "botón cierra antes la cámara y el micrófono.")
    figure(doc, FIG / "fig_sistema_detenido.png",
           "Panel tras pulsar «Detener sistema»",
           "El panel confirma que los modelos se liberaron y que el servidor se apagó; "
           "para volver a usarlo se arranca de nuevo desde la consola.", width_cm=14)

    heading(doc, "4.7. Pruebas y verificación", 2)
    text(doc, "Se actualizaron las pruebas del vídeo y se incorporaron las de ambos "
         "grupos, además de una prueba de integración nueva del audio con el panel. "
         "También se recorrió el panel en un navegador con cámara y micrófono simulados.")
    table(doc, "Resultados de las pruebas automatizadas",
          ["Conjunto", "Qué verifica", "Resultado"],
          [["tests/ (vídeo)", "Rutas de subida retiradas, vídeo en vivo, perfil GTX 1650, parada, alertas", "21 de 21"],
           ["modulos/audio/tests", "Sesiones, PCM, modelo ML (hash y predicciones), integración con el panel", "26 de 26"],
           ["modulos/imagen/tests", "Calidad, preprocesamiento, posturas, ROI, clips, características", "13 de 13"],
           ["Worklet de audio (Node)", "Remuestreo a 16 kHz y cliente sin pista de audio", "5 de 5"]],
          [4.2, 8.8, 3.0],
          "Ejecutadas con python -m pytest y node --test.")

    heading(doc, "4.8. Documentación", 2)
    text(doc, "Se actualizaron el README principal (uso, arquitectura, API, pruebas y "
         "cumplimiento), el registro de decisiones, el diagrama de arquitectura, los "
         "archivos de dependencias, el script de instalación, el diagnóstico del entorno "
         "y el descargador de modelos, y se redactó un README para cada módulo.")

    heading(doc, "5. Limitaciones y trabajo pendiente", 1)
    bullets(doc, [
        "El perfil de la GTX 1650 se validó limitando la memoria de una GPU más potente; "
        "conviene medir en una GTX 1650 física antes de la presentación final.",
        "El audio se probó con señales sintéticas y micrófono simulado; sus umbrales no "
        "están calibrados en una escalera real y sus alertas requieren verificación humana "
        "(ISO, 2023a).",
        "El requisito de pruebas con vídeo pregrabado (RQNF22) queda retirado por pedido del "
        "cliente; un vídeo grabado puede reproducirse como cámara IP en /api/stream.",
        "Las señales de pose del vídeo siguen calibradas con pocas escenas reales.",
    ])

    heading(doc, "6. Conclusiones", 1)
    text(doc, "El sistema quedó operando exclusivamente con vídeo en vivo y cumple los cinco "
         "pedidos del cliente sin reentrenar ningún modelo. Funciona con el perfil de una "
         "GTX 1650 con un consumo de memoria de GPU muy inferior a sus 4 GB; integra el "
         "módulo de audio con todas sus funciones activas y el módulo de imagen en "
         "sustitución del anterior, ambos en carpetas propias que no interfieren con el "
         "vídeo; y dispone de una parada ordenada. Las decisiones se alinearon con la "
         "calidad de producto de ISO/IEC 25010 (ISO, 2023c), con la gestión de IA y de sus "
         "riesgos (ISO, 2023a, 2023b) y con la protección de datos personales exigida por "
         "la Ley N.° 29733 (Congreso de la República del Perú, 2011).")


def referencias(doc):
    page_break(doc)
    heading(doc, "Referencias", 1)
    I = {"i": True}
    N = {}
    refs = [
        [("Congreso de la República del Perú. (2011, 3 de julio). ", N),
         ("Ley N.° 29733, Ley de Protección de Datos Personales", I),
         (". Diario Oficial El Peruano.", N)],
        [("Congreso de la República del Perú. (2023, 5 de julio). ", N),
         ("Ley N.° 31814, Ley que promueve el uso de la inteligencia artificial en favor del "
          "desarrollo económico y social del país", I),
         (". Diario Oficial El Peruano.", N)],
        [("Dirección General de Transparencia, Acceso a la Información Pública y Protección "
          "de Datos Personales. (2020). ", N),
         ("Directiva N.° 01-2020-JUS/DGTAIPD, Tratamiento de datos personales mediante "
          "sistemas de videovigilancia", I),
         (". Ministerio de Justicia y Derechos Humanos.", N)],
        [("International Organization for Standardization. (2018a). ", N),
         ("Systems and software engineering — Life cycle processes — Requirements "
          "engineering", I),
         (" (ISO/IEC/IEEE Standard No. 29148:2018). https://www.iso.org/standard/72089.html", N)],
        [("International Organization for Standardization. (2018b). ", N),
         ("Occupational health and safety management systems — Requirements with guidance "
          "for use", I),
         (" (ISO Standard No. 45001:2018). https://www.iso.org/standard/63787.html", N)],
        [("International Organization for Standardization. (2022a). ", N),
         ("Information security, cybersecurity and privacy protection — Information "
          "security management systems — Requirements", I),
         (" (ISO/IEC Standard No. 27001:2022). https://www.iso.org/standard/27001", N)],
        [("International Organization for Standardization. (2022b). ", N),
         ("Information technology — Artificial intelligence — Artificial intelligence "
          "concepts and terminology", I),
         (" (ISO/IEC Standard No. 22989:2022). https://www.iso.org/standard/74296.html", N)],
        [("International Organization for Standardization. (2023a). ", N),
         ("Information technology — Artificial intelligence — Guidance on risk management", I),
         (" (ISO/IEC Standard No. 23894:2023). https://www.iso.org/standard/77304.html", N)],
        [("International Organization for Standardization. (2023b). ", N),
         ("Information technology — Artificial intelligence — Management system", I),
         (" (ISO/IEC Standard No. 42001:2023). https://www.iso.org/standard/81230.html", N)],
        [("International Organization for Standardization. (2023c). ", N),
         ("Systems and software engineering — Systems and software Quality Requirements "
          "and Evaluation (SQuaRE) — Product quality model", I),
         (" (ISO/IEC Standard No. 25010:2023). https://www.iso.org/standard/78176.html", N)],
        [("Ministerio de Justicia y Derechos Humanos. (2024, 30 de noviembre). ", N),
         ("Decreto Supremo N.° 016-2024-JUS, que aprueba el Reglamento de la Ley N.° 29733, "
          "Ley de Protección de Datos Personales", I),
         (". Diario Oficial El Peruano.", N)],
    ]
    for r in refs:
        reference(doc, r)


def main():
    doc = Document()
    setup(doc)
    _title_style(doc)
    caratula(doc)
    indices(doc)
    cuerpo(doc)
    referencias(doc)
    doc.save(OUT)
    print(f"guardado {OUT}")


if __name__ == "__main__":
    main()
